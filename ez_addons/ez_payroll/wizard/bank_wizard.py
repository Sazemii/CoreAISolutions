from odoo import api,fields,models,tools,_
from odoo.exceptions import ValidationError
import xlwt
from datetime import datetime
from dateutil.relativedelta import relativedelta
import logging
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment,Font
import json,base64
from io import BytesIO
_logger=logging.getLogger(__name__)
DF='%Y-%m-%d'
class Payroll(models.Model):
	_inherit='hr.ph.payroll'
	@api.model
	def create_bank_report_format(self,format,employee_ids=False):
		J=employee_ids;D=self;D.ensure_one();filter=[('payroll_id','=',D.id)]
		if J:filter.append(('employee_id','in',J))
		K=D.env['hr.ph.payslip'].search(filter)
		if not K:raise ValidationError(_('No records found.'))
		L=Workbook();B=L.active;M=Font(bold=True);B.title='Bank Report';H=format.get('title',[]);E=1
		if H:
			for C in H:
				for A in range(len(C)):
					if C[A]=='[name]':C[A]=D.name
			for N in H:
				_logger.debug('Title: %s',N);B.append(N)
				for G in B[E]:G.font=M
				E+=1
		O=format.get('header',[])
		if O:
			B.append(O);A=0
			for G in B[E]:
				try:I=format.get('alignment')[A]
				except:I='left'
				_logger.debug('Align: %s',I);Q=Alignment(horizontal=I,vertical='center');G.font=M;G.alignment=Q;A+=1
			E+=1
		for R in K:
			P=format.get('data',[]);C=[]
			for A in range(len(P)):F=R[P[A]];C.append(F or None)
			B.append(C);E+=1
		F=BytesIO();L.save(F);S=base64.b64encode(F.getvalue());F.close();T=format.get('filename','bank')+' %s.xlsx'%D.name;return S,T
	@api.model
	def get_bank_report_data(self,payslip_ids):
		C=self;A=xlwt.easyxf('font: bold on');F=xlwt.easyxf('font: bold on; align: horiz right');G=xlwt.easyxf(num_format_str='#,##0.00');H=xlwt.easyxf('font: bold on',num_format_str='#,##0.00');D=[[('BANK REPORT',A)],[('Payroll No.:',0),(C.name,0)],[('Date Range:',0),(fields.Date.from_string(C.date_from).strftime('%m/%d/%Y to ')+fields.Date.from_string(C.date_to).strftime('%m/%d/%Y'),0)],[],[('ID Number',A),('Last Name',A),('First Name',A),('Middle Initial',A),('Account No.',A),('Amount',F)]];E=5;I=E+1
		for B in payslip_ids:D.append([(B.employee_id.identification_id or'',0),(B.employee_id.last_name or'',0),(B.employee_id.first_name or'',0),(B.employee_id.middle_name and B.employee_id.middle_name[0]or'',0),(B.employee_id.atm_no or'',0),(B.net_pay or .0,G)]);E+=1
		D.append([('',0),('',0),('',0),('',0),('TOTAL:',A),(xlwt.Formula('SUM(F%s:F%s)'%(I,E)),H)]);return D
	def create_bank_report(A,employee_ids=False):
		D=employee_ids;A.ensure_one();format=A.sudo().env['ir.config_parameter'].get_param('ez_payroll.bank_file_format')
		if format:G=json.loads(format);_logger.debug('New format: %s',G);return A.create_bank_report_format(G,employee_ids=D)
		filter=[('payroll_id','=',A.id)]
		if D:filter.append(('employee_id','in',D))
		H=A.env['hr.ph.payslip'].search(filter)
		if not H:raise ValidationError(_('No records found.'))
		I=xlwt.Workbook();J=I.add_sheet('Bank Report');B=A.get_bank_report_data(H);E=0
		for K in B:
			F=0
			for C in K:
				if C[1]:J.write(E,F,C[0],C[1])
				else:J.write(E,F,C[0])
				F+=1
			E+=1
		B=BytesIO();I.save(B);L=base64.b64encode(B.getvalue());B.close();M='bank-report-%s.xls'%A.name;return L,M
class BankWizard(models.TransientModel):
	_name='ez.bank.wizard';_description='Bank Wizard'
	def get_default_ym(A):B=fields.Date.context_today(A);C=fields.Date.from_string(B);return C.strftime('%Y-%m')
	company_id=fields.Many2one('res.company',string='Company',default=lambda self:self.env.user.company_id.id);payroll_id=fields.Many2one('hr.ph.payroll',string='Payroll',required=True);employee_ids=fields.Many2many('hr.employee',string='Employees');xls_file=fields.Binary(string='Excel Output File',readonly=True);filename=fields.Char()
	def gen_excel(A):A.ensure_one();B=A.payroll_id.create_bank_report(employee_ids=A.employee_ids.ids);A.write({'xls_file':B,'filename':'bank-report-%s.xls'%A.payroll_id.name});return{'name':'Bank Report','view_type':'form','view_mode':'form','res_model':'ez.bank.wizard','res_id':A.id,'type':'ir.actions.act_window','target':'new','context':A._context.copy()}